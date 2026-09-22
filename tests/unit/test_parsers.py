"""Unit tests for document parsers (PDF, DOCX, Markdown, Text)."""

import io

import docx
import pypdf
import pytest

from app.core.exceptions import ParsingError
from app.rag.ingestion.parsers.docx import DocxParser
from app.rag.ingestion.parsers.markdown import MarkdownParser
from app.rag.ingestion.parsers.pdf import PDFParser
from app.rag.ingestion.parsers.text import TextParser


@pytest.mark.asyncio
async def test_text_parser_basic() -> None:
    """Verify TextParser correctly decodes UTF-8 plain text."""
    parser = TextParser()
    raw_content = "Hello world!\nThis is a plain text file."
    parsed = await parser.parse(raw_content.encode("utf-8"), "my_notes.txt")

    assert parsed.title == "My Notes"
    assert "Hello world!" in parsed.content
    assert parsed.metadata["format"] == "text/plain"
    assert parsed.metadata["encoding"] == "utf-8"


@pytest.mark.asyncio
async def test_text_parser_latin1_fallback() -> None:
    """Verify TextParser falls back to latin-1 on non-utf-8 bytes."""
    parser = TextParser()
    raw_content = "Café résumé".encode("latin-1")
    parsed = await parser.parse(raw_content, "cafe.txt")

    assert "Café résumé" in parsed.content
    assert parsed.metadata["encoding"] == "latin-1"


@pytest.mark.asyncio
async def test_markdown_parser_with_frontmatter() -> None:
    """Verify MarkdownParser extracts YAML frontmatter, title, and sections."""
    parser = MarkdownParser()
    md_text = """---
title: System Architecture Guide
author: Team Lead
department: Engineering
---
# System Architecture Guide
This is the introductory overview.

## Data Layer
PostgreSQL is used for relational and vector storage.

### Caching
Redis is used for caching.
"""
    parsed = await parser.parse(md_text.encode("utf-8"), "arch.md")

    assert parsed.title == "System Architecture Guide"
    assert parsed.author == "Team Lead"
    assert parsed.metadata["department"] == "Engineering"
    assert len(parsed.sections) >= 3
    assert any(s["title"] == "Data Layer" for s in parsed.sections)


@pytest.mark.asyncio
async def test_markdown_parser_without_frontmatter() -> None:
    """Verify MarkdownParser infers title from first H1 when frontmatter is absent."""
    parser = MarkdownParser()
    md_text = "# Quickstart Guide\n\nFollow these steps to install the project."
    parsed = await parser.parse(md_text.encode("utf-8"), "quickstart.md")

    assert parsed.title == "Quickstart Guide"
    assert "Follow these steps" in parsed.content


@pytest.mark.asyncio
async def test_pdf_parser_valid() -> None:
    """Verify PDFParser extracts text and page metadata from a valid PDF."""
    # Synthesize a valid multi-page PDF in-memory
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Title": "Test In-Memory PDF", "/Author": "Test Author"})

    pdf_buffer = io.BytesIO()
    writer.write(pdf_buffer)
    pdf_bytes = pdf_buffer.getvalue()

    parser = PDFParser()
    parsed = await parser.parse(pdf_bytes, "sample.pdf")

    assert parsed.metadata["page_count"] == 2
    assert parsed.title == "Test In-Memory PDF"
    assert parsed.author == "Test Author"


@pytest.mark.asyncio
async def test_pdf_parser_corrupted_bytes_raises() -> None:
    """Verify PDFParser raises ParsingError when provided with corrupted PDF bytes."""
    parser = PDFParser()
    with pytest.raises(ParsingError) as exc_info:
        await parser.parse(b"%PDF-corrupted-and-truncated", "corrupted.pdf")
    assert "Corrupted or invalid PDF format" in str(exc_info.value)


@pytest.mark.asyncio
async def test_docx_parser_valid() -> None:
    """Verify DocxParser extracts text, headings, tables, and core properties."""
    doc = docx.Document()
    doc.core_properties.title = "Project Proposal"
    doc.core_properties.author = "Jane Doe"

    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph("This proposal outlines the timeline and budget.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Cost"
    table.cell(1, 0).text = "Servers"
    table.cell(1, 1).text = "$500"

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()

    parser = DocxParser()
    parsed = await parser.parse(docx_bytes, "proposal.docx")

    assert parsed.title == "Project Proposal"
    assert parsed.author == "Jane Doe"
    assert "This proposal outlines" in parsed.content
    assert "Servers | $500" in parsed.content
    assert parsed.metadata["table_count"] == 1


@pytest.mark.asyncio
async def test_docx_parser_corrupted_bytes_raises() -> None:
    """Verify DocxParser raises ParsingError on malformed DOCX archive."""
    parser = DocxParser()
    with pytest.raises(ParsingError) as exc_info:
        await parser.parse(b"PK\x03\x04-corrupted-zip-stream", "bad.docx")
    assert "Failed to open or parse DOCX" in str(exc_info.value)
