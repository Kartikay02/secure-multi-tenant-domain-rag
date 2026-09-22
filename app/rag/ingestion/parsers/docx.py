"""DOCX (Microsoft Word) document parser using python-docx."""

import io
from pathlib import Path

import docx

from app.core.exceptions import ParsingError
from app.rag.ingestion.parsers.base import DocumentParserProtocol, ParsedDocument


class DocxParser(DocumentParserProtocol):
    """Extracts text, headings, tables, and document properties from DOCX files."""

    def supported_extensions(self) -> set[str]:
        return {".docx"}

    async def parse(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        """Parse DOCX bytes into structured sections and metadata."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as exc:
            raise ParsingError(
                filename=filename, reason=f"Failed to open or parse DOCX: {exc}"
            ) from exc

        sections: list[dict[str, str]] = []
        current_section_title = "Document Body"
        current_section_lines: list[str] = []

        # Extract paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            if style_name.startswith("Heading"):
                if current_section_lines:
                    sections.append(
                        {
                            "title": current_section_title,
                            "content": "\n".join(current_section_lines),
                        }
                    )
                    current_section_lines = []
                current_section_title = text
            else:
                current_section_lines.append(text)

        # Extract tables as formatted tabular text
        for table_idx, table in enumerate(doc.tables, start=1):
            table_lines: list[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_lines.append(" | ".join(row_cells))
            if table_lines:
                current_section_lines.append(f"\n[Table {table_idx}]\n" + "\n".join(table_lines))

        if current_section_lines:
            sections.append(
                {
                    "title": current_section_title,
                    "content": "\n".join(current_section_lines),
                }
            )

        full_content = "\n\n".join(s["content"] for s in sections)

        # Extract Core Properties
        core_props = doc.core_properties
        title = (
            core_props.title if core_props.title else Path(filename).stem.replace("_", " ").title()
        )
        author = core_props.author if core_props.author else None

        return ParsedDocument(
            content=full_content,
            title=title,
            author=author,
            sections=sections,
            metadata={
                "format": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "created": str(core_props.created) if core_props.created else "",
                "modified": str(core_props.modified) if core_props.modified else "",
            },
        )
